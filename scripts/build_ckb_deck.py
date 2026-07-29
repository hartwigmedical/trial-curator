#!/usr/bin/env python3
"""Build the CKB presentation .pptx from the HTML deck + the official Hartwig template.

Two steps, both repeatable — edit a slide in the HTML and re-run:

  1. RENDER  each content slide of `docs/v2_agentic_presentation_ckb.html` to a PNG with headless
     Chrome, at the exact canvas size (1280x638 CSS px, device scale 2 -> 2560x1276).
  2. ASSEMBLE `docs/v2_agentic_presentation_ckb.pptx` from `docs/TemplatePPTX.pptx`:
       - the title slide, the two act dividers and the closing slide stay NATIVE template slides
         (real branding, editable text);
       - every other slide is the rendered PNG, placed full width at (0, 0) sized
         13.3333in x 6.6421in — i.e. stopping exactly where the master's branded bottom bar starts,
         so that bar stays visible and authentic.

The HTML marks its native slides with `data-native="..."`; they are rendered as stand-ins on screen and
skipped here.

  python3 scripts/build_ckb_deck.py              # render + assemble
  python3 scripts/build_ckb_deck.py --no-render  # assemble from the PNGs already in build/
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

REPO = Path(__file__).resolve().parents[1]
HTML = REPO / "docs" / "v2_agentic_presentation_ckb.html"
TEMPLATE = REPO / "docs" / "TemplatePPTX.pptx"
OUT_PPTX = REPO / "docs" / "v2_agentic_presentation_ckb.pptx"
BUILD = REPO / "docs" / "build" / "ckb_slides"

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
CANVAS_W, CANVAS_H = 1280, 638          # CSS px; must match the .stage size in the HTML
SCALE = 2                               # device pixel ratio -> 2560 x 1276 PNGs

# Placement: full slide width, stopping at the master's branded bottom bar (y = 6,073,512 EMU).
IMG_LEFT, IMG_TOP = Emu(0), Emu(0)
IMG_W, IMG_H = Emu(12192000), Emu(6073512)

BLANK_LAYOUT = "Leeg"                   # the template's empty layout (still shows master branding)

# Native template slides to reuse, by their index in TemplatePPTX.pptx (0-based).
TPL_TITLE, TPL_DIV_A, TPL_DIV_B, TPL_CLOSE = 0, 3, 4, 8

TITLE_TEXT = "Automating trial curation"
SUBTITLE_TEXT = "An agentic pipeline for ~2,000 free-text cancer trials"
SPEAKER_TEXT = "Junran Cao  ·  Hartwig Medical Foundation  ·  30 July 2026"
DIV_A_TEXT = "The pipeline, end to end"
DIV_B_TEXT = "What it produces, and what we learned"


# ----------------------------------------------------------------------------- step 1: render
def slide_plan() -> list[dict]:
    """Read the HTML and return one entry per slide, in order, flagging the native ones."""
    html = HTML.read_text()
    plan = []
    for n, tag in enumerate(re.findall(r'<section class="slide[^"]*"([^>]*)>', html), start=1):
        native = re.search(r'data-native="([^"]+)"', tag)
        plan.append({"n": n, "native": native.group(1) if native else None})
    if not plan:
        sys.exit("no <section class=\"slide\"> found in the HTML")
    return plan


def render(plan: list[dict]) -> None:
    if not Path(CHROME).exists():
        sys.exit(f"Chrome not found at {CHROME}")
    if BUILD.exists():
        shutil.rmtree(BUILD)
    BUILD.mkdir(parents=True)
    todo = [s for s in plan if not s["native"]]
    print(f"rendering {len(todo)} content slide(s) at {CANVAS_W * SCALE}x{CANVAS_H * SCALE} …")
    for s in todo:
        out = BUILD / f"slide{s['n']:02d}.png"
        url = f"file://{HTML}?export=1#{s['n']}"
        subprocess.run(
            [CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars",
             f"--force-device-scale-factor={SCALE}", f"--window-size={CANVAS_W},{CANVAS_H}",
             "--virtual-time-budget=4000", f"--screenshot={out}", url],
            check=True, capture_output=True,
        )
        if not out.exists():
            sys.exit(f"render failed for slide {s['n']}")
        print(f"  slide {s['n']:>2}  ->  {out.name}  ({out.stat().st_size // 1024} KB)")


# ----------------------------------------------------------------------------- step 2: assemble
def set_placeholder_text(shape, text: str) -> None:
    """Replace a shape's text, keeping the first run's formatting (so brand styling survives)."""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    if first.runs:
        first.runs[0].text = text
        for extra in first.runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        first.add_run().text = text
    for p in list(tf.paragraphs[1:]):          # drop the template's filler paragraphs
        p._p.getparent().remove(p._p)


def find_text_shape(slide, contains: str):
    for sh in slide.shapes:
        if sh.has_text_frame and contains.lower() in sh.text_frame.text.lower():
            return sh
    return None


def assemble(plan: list[dict]) -> None:
    prs = Presentation(str(TEMPLATE))
    sld_id_lst = prs.slides._sldIdLst
    original = list(sld_id_lst)

    # --- fill the native slides we keep -------------------------------------------------
    title = prs.slides[TPL_TITLE]
    for sh in title.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip().lower()
        if t == "title":
            set_placeholder_text(sh, TITLE_TEXT)
        elif t == "subtitle":
            set_placeholder_text(sh, SUBTITLE_TEXT)
        elif t == "speaker":
            set_placeholder_text(sh, SPEAKER_TEXT)

    for idx, text in ((TPL_DIV_A, DIV_A_TEXT), (TPL_DIV_B, DIV_B_TEXT)):
        sh = find_text_shape(prs.slides[idx], "divider")
        if sh is None:
            sys.exit(f"template slide {idx + 1} has no divider text placeholder")
        set_placeholder_text(sh, text)

    keep = {TPL_TITLE: original[TPL_TITLE], TPL_DIV_A: original[TPL_DIV_A],
            TPL_DIV_B: original[TPL_DIV_B], TPL_CLOSE: original[TPL_CLOSE]}

    # --- one blank slide per rendered PNG ------------------------------------------------
    blank = next((l for l in prs.slide_layouts if l.name == BLANK_LAYOUT), None)
    if blank is None:
        sys.exit(f"template has no '{BLANK_LAYOUT}' layout")
    image_ids: dict[int, object] = {}
    for s in plan:
        if s["native"]:
            continue
        png = BUILD / f"slide{s['n']:02d}.png"
        if not png.exists():
            sys.exit(f"missing render for slide {s['n']}: {png} (drop --no-render)")
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), IMG_LEFT, IMG_TOP, width=IMG_W, height=IMG_H)
        image_ids[s["n"]] = sld_id_lst[-1]

    # --- order: native title / dividers / close in their outline positions ---------------
    native_slot = {"title": keep[TPL_TITLE], "close": keep[TPL_CLOSE]}
    dividers = [keep[TPL_DIV_A], keep[TPL_DIV_B]]
    ordered = []
    for s in plan:
        if s["native"] == "divider":
            ordered.append(dividers.pop(0))
        elif s["native"]:
            ordered.append(native_slot[s["native"]])
        else:
            ordered.append(image_ids[s["n"]])

    for el in list(sld_id_lst):                 # drop every template slide we did not keep
        if el not in ordered:
            sld_id_lst.remove(el)
    for el in ordered:                          # then re-append in outline order
        sld_id_lst.remove(el)
        sld_id_lst.append(el)

    prs.save(str(OUT_PPTX))
    natives = sum(1 for s in plan if s["native"])
    print(f"\nwrote {OUT_PPTX.relative_to(REPO)}"
          f"\n  {len(plan)} slides — {natives} native template slides, {len(plan) - natives} rendered"
          f"\n  images placed at (0, 0) 13.3333in x 6.6421in — the master's bottom bar stays visible")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--no-render", action="store_true", help="reuse the PNGs already in docs/build/ckb_slides")
    args = ap.parse_args()

    plan = slide_plan()
    print(f"{len(plan)} slides in the HTML — "
          f"native: {[s['n'] for s in plan if s['native']]}")
    if not args.no_render:
        render(plan)
    assemble(plan)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
