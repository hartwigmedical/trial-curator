from __future__ import annotations

import ast
import logging
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Set, Tuple

import pandas as pd
from openpyxl import load_workbook

from aus_trial_universe.ctgov.utils.general.text_normalisation import (
    clean_cell_str,
    is_effectively_empty,
)

logger = logging.getLogger(__name__)

DELIMITER = " | "
NULL_SENTINELS = {"[None]", "NOT([None])"}


def _normalize_string(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _clean_mapped_value(value: str) -> str:
    stripped = clean_cell_str(value)
    if stripped in NULL_SENTINELS:
        return ""
    return stripped


def _dedupe_preserve_order(values: Sequence[str]) -> List[str]:
    seen: Set[str] = set()
    output: List[str] = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            output.append(value)
    return output


def join_condition_terms(values: Sequence[str]) -> str:
    return DELIMITER.join(
        _dedupe_preserve_order(
            [clean_cell_str(v) for v in values if not is_effectively_empty(v)]
        )
    )


def parse_condition_terms(raw_value: object) -> List[str]:
    if raw_value is None:
        return []

    if isinstance(raw_value, list):
        return [clean_cell_str(item) for item in raw_value if not is_effectively_empty(item)]

    text = str(raw_value).strip()
    if not text:
        return []

    try:
        parsed = ast.literal_eval(text)
    except (SyntaxError, ValueError):
        return [clean_cell_str(text)] if not is_effectively_empty(text) else []

    if isinstance(parsed, list):
        return [clean_cell_str(item) for item in parsed if not is_effectively_empty(item)]

    return [clean_cell_str(text)] if not is_effectively_empty(text) else []


def load_conditions_mapping(
    mapping_excel: Path,
    *,
    sheet_name: Optional[str] = None,
    lookup_column_name: str = "Conditions_lookup",
    value_column_name: str = "Oncotree_curation",
) -> Dict[str, str]:
    wb = load_workbook(mapping_excel, read_only=True, data_only=True)
    target_sheet_name = sheet_name or wb.sheetnames[0]
    ws = wb[target_sheet_name]

    header_to_col: Dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is not None:
            header_to_col[str(header).strip()] = col_idx

    if lookup_column_name not in header_to_col:
        raise ValueError(f"Could not find header '{lookup_column_name}' in worksheet '{ws.title}'")
    if value_column_name not in header_to_col:
        raise ValueError(f"Could not find header '{value_column_name}' in worksheet '{ws.title}'")

    lookup_col_idx = header_to_col[lookup_column_name]
    value_col_idx = header_to_col[value_column_name]

    mapping: Dict[str, str] = {}
    for row_idx in range(2, ws.max_row + 1):
        lookup_value = _normalize_string(ws.cell(row=row_idx, column=lookup_col_idx).value)
        curated_value = _normalize_string(ws.cell(row=row_idx, column=value_col_idx).value)
        if not lookup_value:
            continue
        mapping[lookup_value] = curated_value

    return mapping


def map_condition_terms(
    raw_value: object,
    conditions_mapping: Dict[str, str],
    *,
    missing_condition_terms: Optional[Set[str]] = None,
) -> Tuple[str, str]:
    original_terms = parse_condition_terms(raw_value)
    mapped_terms: List[str] = []

    for term in original_terms:
        mapped_value = _clean_mapped_value(conditions_mapping.get(term, ""))
        if mapped_value:
            mapped_terms.append(mapped_value)
        elif missing_condition_terms is not None and term:
            missing_condition_terms.add(term)

    return join_condition_terms(original_terms), join_condition_terms(mapped_terms)


def load_trial_conditions_lookup(
    ctgov_extractions_csv: Path,
    *,
    nct_id_column: str = "nctId",
    conditions_column: str = "conditions",
) -> Dict[str, object]:
    df = pd.read_csv(
        ctgov_extractions_csv,
        keep_default_na=False,
        na_values=[],
        dtype=str,
    )

    missing = [c for c in (nct_id_column, conditions_column) if c not in df.columns]
    if missing:
        raise ValueError(
            f"ctgov extractions CSV missing required columns: {missing}. Found: {list(df.columns)}"
        )

    lookup: Dict[str, object] = {}
    for _, row in df.iterrows():
        nct_id = _normalize_string(row.get(nct_id_column)).upper()
        if not nct_id:
            continue
        lookup[nct_id] = row.get(conditions_column, "")

    return lookup


def build_trial_conditions_enrichment_lookup(
    ctgov_extractions_csv: Path,
    conditions_mapping_excel: Path,
    *,
    nct_id_column: str = "nctId",
    conditions_column: str = "conditions",
    mapping_sheet_name: Optional[str] = None,
    mapping_lookup_column_name: str = "Conditions_lookup",
    mapping_value_column_name: str = "Oncotree_curation",
    missing_condition_terms: Optional[Set[str]] = None,
) -> Dict[str, Dict[str, str]]:
    trial_conditions_lookup = load_trial_conditions_lookup(
        ctgov_extractions_csv,
        nct_id_column=nct_id_column,
        conditions_column=conditions_column,
    )
    conditions_mapping = load_conditions_mapping(
        conditions_mapping_excel,
        sheet_name=mapping_sheet_name,
        lookup_column_name=mapping_lookup_column_name,
        value_column_name=mapping_value_column_name,
    )

    enriched: Dict[str, Dict[str, str]] = {}
    for nct_id, raw_conditions in trial_conditions_lookup.items():
        conditions_original, conditions_oncotree_curation = map_condition_terms(
            raw_conditions,
            conditions_mapping,
            missing_condition_terms=missing_condition_terms,
        )
        enriched[nct_id] = {
            "conditions_original": conditions_original,
            "conditions_oncotree_curation": conditions_oncotree_curation,
        }

    return enriched
