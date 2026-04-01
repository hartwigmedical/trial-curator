from __future__ import annotations

import argparse
import ast
import logging
import re
from collections import defaultdict
from copy import copy
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

TARGET_CRITERIA = {
    "GeneAlterationCriterion": ("GeneAlteration", "gene_alteration_curation", False),
    "PrimaryTumorCriterion": ("PrimaryTumor", "Oncotree_curation", False),
}

DELIMITER = " | "
NCT_ID_RE = re.compile(r"(NCT\d{8})", re.IGNORECASE)
NOT_WRAPPER_RE = re.compile(r"^NOT\((.*)\)$")
LEVEL_COLUMN_RE = re.compile(r"^level_(\d+)$")

CLEANABLE_LABELS = {"PrimaryTumor", "CuratedConditions"}
NULL_SENTINELS = {"[None]", "NOT([None])"}

PAN_CANCER = "Pan-Cancer"

DEFAULT_SHEET_NAME = "general"
DEFAULT_NCTID_COLUMN = "nctId"
DEFAULT_CONDITIONS_COLUMN = "conditions"
DEFAULT_STEP1_COLUMN = "CancerType_step1"
DEFAULT_STEP2_COLUMN = "CancerType_step2"
DEFAULT_STEP3_COLUMN = "CancerType_step3"


# ---------------------------------------------------------------------------
# Generic worksheet helpers
# ---------------------------------------------------------------------------


def _find_header_column(ws, header_name: str) -> int:
    for col_idx in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=col_idx).value == header_name:
            return col_idx
    raise ValueError(f"Could not find header '{header_name}' in worksheet '{ws.title}'")


def _copy_cell_style(src, dst) -> None:
    if src.has_style:
        dst._style = copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.font is not None:
        dst.font = copy(src.font)
    if src.fill is not None:
        dst.fill = copy(src.fill)
    if src.border is not None:
        dst.border = copy(src.border)
    if src.alignment is not None:
        dst.alignment = copy(src.alignment)
    if src.protection is not None:
        dst.protection = copy(src.protection)


def _normalize_string(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _prepare_output_columns(ws, column_names: Sequence[str]) -> Dict[str, int]:
    col_map: Dict[str, int] = {}

    if ws.max_column < 1:
        raise ValueError(f"Worksheet '{ws.title}' has no columns")

    start_col = ws.max_column + 1
    header_style_template = ws.cell(row=1, column=ws.max_column)
    body_style_template = ws.cell(row=2, column=ws.max_column) if ws.max_row >= 2 else None
    source_width = ws.column_dimensions[get_column_letter(ws.max_column)].width

    for offset, column_name in enumerate(column_names):
        col_idx = start_col + offset
        col_map[column_name] = col_idx

        header_cell = ws.cell(row=1, column=col_idx)
        header_cell.value = column_name
        _copy_cell_style(header_style_template, header_cell)

        if body_style_template is not None:
            for row_idx in range(2, ws.max_row + 1):
                _copy_cell_style(body_style_template, ws.cell(row=row_idx, column=col_idx))

        ws.column_dimensions[get_column_letter(col_idx)].width = max(
            (source_width or 13) * 2.5,
            len(column_name) + 2,
            30,
        )

    return col_map


# ---------------------------------------------------------------------------
# Criteria extraction from curated NCT*.py files
# ---------------------------------------------------------------------------
def _split_top_level_pipe(expr: str) -> list[str]:
    """
    Split on top-level ' | ' only, respecting nesting inside (), [], {}.
    """
    if not expr:
        return []

    parts: list[str] = []
    buf: list[str] = []
    paren = 0
    bracket = 0
    brace = 0

    i = 0
    n = len(expr)

    while i < n:
        ch = expr[i]

        if ch == "(":
            paren += 1
            buf.append(ch)
            i += 1
            continue
        if ch == ")":
            paren = max(paren - 1, 0)
            buf.append(ch)
            i += 1
            continue
        if ch == "[":
            bracket += 1
            buf.append(ch)
            i += 1
            continue
        if ch == "]":
            bracket = max(bracket - 1, 0)
            buf.append(ch)
            i += 1
            continue
        if ch == "{":
            brace += 1
            buf.append(ch)
            i += 1
            continue
        if ch == "}":
            brace = max(brace - 1, 0)
            buf.append(ch)
            i += 1
            continue

        if (
            paren == 0
            and bracket == 0
            and brace == 0
            and expr[i:i+3] == " | "
        ):
            part = "".join(buf).strip()
            if part:
                parts.append(part)
            buf = []
            i += 3
            continue

        buf.append(ch)
        i += 1

    tail = "".join(buf).strip()
    if tail:
        parts.append(tail)

    return parts


def _is_balanced_wrapped_not(term: str) -> bool:
    """
    True only if the whole term is exactly NOT(<content>) at top level.
    """
    if not term:
        return False

    term = term.strip()
    if not term.startswith("NOT(") or not term.endswith(")"):
        return False

    inner = term[4:-1]
    if not inner:
        return False

    paren = 0
    bracket = 0
    brace = 0

    for i, ch in enumerate(term):
        if ch == "(":
            paren += 1
        elif ch == ")":
            paren -= 1
            if paren < 0:
                return False
            # before the very last char, paren should not drop to 0
            if paren == 0 and i != len(term) - 1:
                return False
        elif ch == "[":
            bracket += 1
        elif ch == "]":
            bracket -= 1
            if bracket < 0:
                return False
        elif ch == "{":
            brace += 1
        elif ch == "}":
            brace -= 1
            if brace < 0:
                return False

    return paren == 0 and bracket == 0 and brace == 0


def _unwrap_gene_alt_not(term: str) -> str:
    term = term.strip()
    return term[4:-1].strip()


def _normalize_gene_alt_unit(term: str) -> str:
    """
    Normalization for exact top-level contradiction matching only.
    Keeps logic unchanged; only normalizes whitespace.
    """
    return re.sub(r"\s+", " ", term.strip())


def _build_gene_alteration_postprocessed(raw_value: str) -> tuple[str, str]:
    """
    Returns:
        (GeneAlteration-Inclusive, GeneAlteration-Exclusive)

    Rules:
    - split original GeneAlteration on top-level ' | '
    - positive top-level terms -> Inclusive
    - negative top-level NOT(...) terms -> Exclusive
    - contradiction cleanup on exact top-level units only:
          X  versus  NOT(X)
      removes both
    - Exclusive is emitted in compressed form:
          NOT(A | B | C)
    """
    if raw_value is None:
        return "", ""

    raw_value = str(raw_value).strip()
    if not raw_value:
        return "", ""

    top_level_terms = _split_top_level_pipe(raw_value)

    positive_terms: list[str] = []
    negative_inner_terms: list[str] = []

    for term in top_level_terms:
        term = term.strip()
        if not term:
            continue

        if _is_balanced_wrapped_not(term):
            negative_inner_terms.append(_unwrap_gene_alt_not(term))
        else:
            positive_terms.append(term)

    # exact top-level contradiction cleanup
    positive_norms = {_normalize_gene_alt_unit(t) for t in positive_terms}
    negative_norms = {_normalize_gene_alt_unit(t) for t in negative_inner_terms}
    contradictions = positive_norms & negative_norms

    if contradictions:
        positive_terms = [
            t for t in positive_terms
            if _normalize_gene_alt_unit(t) not in contradictions
        ]
        negative_inner_terms = [
            t for t in negative_inner_terms
            if _normalize_gene_alt_unit(t) not in contradictions
        ]

    inclusive_value = " | ".join(positive_terms) if positive_terms else ""

    if negative_inner_terms:
        exclusive_value = f"NOT({' | '.join(negative_inner_terms)})"
    else:
        exclusive_value = ""

    return inclusive_value, exclusive_value


def _iter_input_py_files(curated_dir: Path) -> Iterable[Path]:
    if curated_dir.is_file():
        yield curated_dir
        return
    yield from sorted(p for p in curated_dir.glob("NCT*.py") if p.is_file())


def _extract_nct_id_from_path(py_path: Path) -> str:
    match = NCT_ID_RE.search(py_path.stem)
    if not match:
        raise ValueError(f"Could not extract NCT ID from filename: {py_path.name}")
    return match.group(1).upper()


def _call_name(node: ast.AST) -> Optional[str]:
    if isinstance(node, ast.Name):
        return node.id
    if isinstance(node, ast.Attribute):
        return node.attr
    return None


def _get_call_argument(call: ast.Call, key: str) -> Optional[ast.AST]:
    for kw in call.keywords:
        if kw.arg == key:
            return kw.value
    return None


def _get_string_constant(node: Optional[ast.AST]) -> Optional[str]:
    if isinstance(node, ast.Constant) and isinstance(node.value, str):
        return node.value
    return None


def _split_emitted_values(raw_value: str, split_on_delimiter: bool) -> List[str]:
    if split_on_delimiter:
        return [part.strip() for part in raw_value.split(DELIMITER) if part.strip()]
    stripped = raw_value.strip()
    return [stripped] if stripped else []


def _apply_not_wrapper(value: str, is_negated: bool) -> str:
    if not is_negated:
        return value
    if value.startswith("NOT("):
        return value
    return f"NOT({value})"


def _clean_value(criterion_label: str, value: str) -> Optional[str]:
    stripped = value.strip()
    if criterion_label in CLEANABLE_LABELS and stripped in NULL_SENTINELS:
        return None
    return stripped


def _finalize_value(criterion_label: str, value: str, *, is_negated: bool) -> Optional[str]:
    wrapped = _apply_not_wrapper(value.strip(), is_negated=is_negated)
    return _clean_value(criterion_label, wrapped)


def _add_extracted_value(
    bucket: Dict[str, List[str]],
    criterion_label: str,
    raw_value: str,
    *,
    is_negated: bool,
    split_on_delimiter: bool,
) -> None:
    parts = _split_emitted_values(raw_value, split_on_delimiter=split_on_delimiter)
    if not parts:
        return

    for part in parts:
        final_value = _finalize_value(
            criterion_label,
            part,
            is_negated=is_negated,
        )
        if final_value is not None:
            bucket.setdefault(criterion_label, []).append(final_value)


def _walk_ast(node: ast.AST, bucket: Dict[str, List[str]], *, negation_depth: int = 0) -> None:
    if isinstance(node, ast.Call):
        func_name = _call_name(node.func)

        if func_name == "NotCriterion":
            child = _get_call_argument(node, "criterion")
            if child is None and node.args:
                child = node.args[0]
            if child is not None:
                _walk_ast(child, bucket, negation_depth=negation_depth + 1)
            return

        target_spec = TARGET_CRITERIA.get(func_name)
        if target_spec is not None:
            criterion_label, value_arg_name, split_on_delimiter = target_spec
            raw_value = _get_string_constant(_get_call_argument(node, value_arg_name))
            if raw_value:
                _add_extracted_value(
                    bucket,
                    criterion_label,
                    raw_value,
                    is_negated=negation_depth > 0,
                    split_on_delimiter=split_on_delimiter,
                )

    for child in ast.iter_child_nodes(node):
        _walk_ast(child, bucket, negation_depth=negation_depth)


def _collapse_primary_tumor_contradictions(values: Sequence[str]) -> List[str]:
    value_set = set(values)
    collapsed: Set[str] = set()

    for value in value_set:
        if value.startswith("NOT(") and value.endswith(")"):
            positive_value = value[4:-1]
            if positive_value in value_set:
                continue
        collapsed.add(value)

    return sorted(collapsed)


def _dedupe_and_sort(criterion_label: str, values: Sequence[str]) -> List[str]:
    deduped = {value for value in values if value}
    if criterion_label == "PrimaryTumor":
        return _collapse_primary_tumor_contradictions(deduped)
    return sorted(deduped)


def _join_cleaned_values(criterion_label: str, values: Sequence[str]) -> str:
    return DELIMITER.join(_dedupe_and_sort(criterion_label, values))


def extract_criteria_from_py(py_path: Path) -> Dict[str, List[str]]:
    tree = ast.parse(py_path.read_text(encoding="utf-8"), filename=str(py_path))
    extracted: Dict[str, List[str]] = {}
    _walk_ast(tree, extracted)
    return {
        criterion: _dedupe_and_sort(criterion, values)
        for criterion, values in extracted.items()
        if values
    }


def build_trial_criteria_lookup(curated_dir: Path) -> Dict[str, Dict[str, List[str]]]:
    lookup: Dict[str, Dict[str, List[str]]] = {}
    seen_paths: Dict[str, Path] = {}

    for py_path in _iter_input_py_files(curated_dir):
        nct_id = _extract_nct_id_from_path(py_path)
        if nct_id in lookup:
            raise ValueError(
                f"Duplicate curated .py files for {nct_id}: {seen_paths[nct_id]} and {py_path}"
            )
        seen_paths[nct_id] = py_path
        lookup[nct_id] = extract_criteria_from_py(py_path)

    return lookup


# ---------------------------------------------------------------------------
# Conditions mapping
# ---------------------------------------------------------------------------


def _parse_condition_terms(raw_value: object) -> List[str]:
    if raw_value is None:
        return []

    if isinstance(raw_value, list):
        return [str(item).strip() for item in raw_value if str(item).strip()]

    text = str(raw_value).strip()
    if not text:
        return []

    try:
        parsed = ast.literal_eval(text)
    except (SyntaxError, ValueError):
        return [text]

    if isinstance(parsed, list):
        return [str(item).strip() for item in parsed if str(item).strip()]

    return [text]


def _map_condition_terms(
    raw_value: object,
    conditions_mapping: Dict[str, str],
    *,
    missing_condition_terms: Set[str],
) -> str:
    mapped_terms: List[str] = []

    for term in _parse_condition_terms(raw_value):
        mapped_term = _normalize_string(conditions_mapping.get(term, ""))
        if mapped_term:
            cleaned = _clean_value("CuratedConditions", mapped_term)
            if cleaned is not None:
                mapped_terms.append(cleaned)
        else:
            missing_condition_terms.add(term)

    return _join_cleaned_values("CuratedConditions", mapped_terms)


def load_conditions_mapping(
    mapping_excel: Path,
    *,
    sheet_name: Optional[str] = None,
    lookup_column_name: str = "Conditions_lookup",
    value_column_name: str = "Oncotree_curation",
) -> Dict[str, str]:
    wb = load_workbook(mapping_excel, read_only=True, data_only=True)
    target_sheet_name = sheet_name or wb.sheetnames[0]
    ws = wb[target_sheet_name]

    lookup_col_idx = _find_header_column(ws, lookup_column_name)
    value_col_idx = _find_header_column(ws, value_column_name)

    mapping: Dict[str, str] = {}
    for row_idx in range(2, ws.max_row + 1):
        lookup_value = _normalize_string(ws.cell(row=row_idx, column=lookup_col_idx).value)
        curated_value = _normalize_string(ws.cell(row=row_idx, column=value_col_idx).value)

        if not lookup_value:
            continue
        mapping[lookup_value] = curated_value

    return mapping


# ---------------------------------------------------------------------------
# CancerType step 1
# ---------------------------------------------------------------------------


def _split_terms(raw_value: object) -> List[str]:
    text = _normalize_string(raw_value)
    if not text:
        return []
    return [part.strip() for part in text.split(DELIMITER) if part.strip()]


def _unwrap_not(term: str) -> Optional[str]:
    match = NOT_WRAPPER_RE.fullmatch(term.strip())
    if not match:
        return None
    inner = match.group(1).strip()
    return inner if inner else None


def _is_not_term(term: str) -> bool:
    return _unwrap_not(term) is not None


def _dedupe_preserve_order(values: Sequence[str]) -> List[str]:
    seen: Set[str] = set()
    output: List[str] = []
    for value in values:
        if value not in seen:
            seen.add(value)
            output.append(value)
    return output


def _drop_negated_when_positive_exists(values: Sequence[str]) -> List[str]:
    value_set = set(values)
    output: List[str] = []

    for value in values:
        positive = _unwrap_not(value)
        if positive is not None and positive in value_set:
            continue
        output.append(value)

    return output


def _all_terms_are_negative(values: Sequence[str]) -> bool:
    return bool(values) and all(_is_not_term(value) for value in values)


def _join_terms(values: Sequence[str]) -> str:
    return DELIMITER.join(values)


def _resolve_cancer_type_step1(primary_tumor_raw: object, curated_conditions_raw: object) -> str:
    primary_terms = _split_terms(primary_tumor_raw)
    curated_terms = _split_terms(curated_conditions_raw)

    if not primary_terms and not curated_terms:
        return PAN_CANCER

    if not curated_terms:
        primary_terms = _dedupe_preserve_order(primary_terms)
        if _all_terms_are_negative(primary_terms):
            return PAN_CANCER
        return _join_terms(primary_terms)

    if not primary_terms:
        curated_terms = _dedupe_preserve_order(curated_terms)
        return _join_terms(curated_terms)

    combined_terms = primary_terms + curated_terms
    combined_terms = _dedupe_preserve_order(combined_terms)
    combined_terms = _drop_negated_when_positive_exists(combined_terms)

    if not combined_terms:
        return PAN_CANCER

    if _all_terms_are_negative(combined_terms):
        return PAN_CANCER

    return _join_terms(combined_terms)


# ---------------------------------------------------------------------------
# CancerType step 2
# ---------------------------------------------------------------------------

class OncoTreeHierarchy:
    def __init__(
        self,
        parents_by_term: Dict[str, Set[str]],
        children_by_term: Dict[str, Set[str]],
        levels_by_term: Dict[str, int],
    ) -> None:
        self.parents_by_term = parents_by_term
        self.children_by_term = children_by_term
        self.levels_by_term = levels_by_term

    def has_term(self, term: str) -> bool:
        return term in self.levels_by_term

    def descendants_of(self, term: str) -> Set[str]:
        visited: Set[str] = set()
        stack = list(self.children_by_term.get(term, set()))

        while stack:
            current = stack.pop()
            if current in visited:
                continue
            visited.add(current)
            stack.extend(self.children_by_term.get(current, set()))

        return visited

    def is_more_granular(self, candidate: str, reference: str) -> bool:
        return candidate in self.descendants_of(reference)


def load_oncotree_hierarchy(oncotree_csv: Path) -> OncoTreeHierarchy:
    df = pd.read_csv(oncotree_csv, dtype=str).fillna("")

    level_columns: List[Tuple[int, str]] = []
    for column in df.columns:
        match = LEVEL_COLUMN_RE.fullmatch(column)
        if match:
            level_columns.append((int(match.group(1)), column))

    if not level_columns:
        raise ValueError(
            f"No level_* columns found in OncoTree CSV '{oncotree_csv}'. "
            f"Found columns: {list(df.columns)}"
        )

    level_columns.sort()
    ordered_level_cols = [column for _, column in level_columns]

    parents_by_term: Dict[str, Set[str]] = defaultdict(set)
    children_by_term: Dict[str, Set[str]] = defaultdict(set)
    levels_by_term: Dict[str, int] = {}

    for _, row in df.iterrows():
        branch_terms: List[Tuple[str, int]] = []
        for level_idx, column_name in enumerate(ordered_level_cols, start=1):
            term = str(row[column_name]).strip()
            if not term:
                continue
            branch_terms.append((term, level_idx))
            levels_by_term[term] = max(levels_by_term.get(term, 0), level_idx)

        for i in range(len(branch_terms) - 1):
            parent_term, _ = branch_terms[i]
            child_term, _ = branch_terms[i + 1]
            parents_by_term[child_term].add(parent_term)
            children_by_term[parent_term].add(child_term)

        for term, _ in branch_terms:
            parents_by_term.setdefault(term, set())
            children_by_term.setdefault(term, set())

    return OncoTreeHierarchy(
        parents_by_term=parents_by_term,
        children_by_term=children_by_term,
        levels_by_term=levels_by_term,
    )


def _extract_positive_terms(raw_value: object) -> List[str]:
    positive_terms: List[str] = []
    for term in _split_terms(raw_value):
        if not _is_not_term(term):
            positive_terms.append(term)
    return _dedupe_preserve_order(positive_terms)


def _extract_negative_terms(raw_value: object) -> List[str]:
    negative_terms: List[str] = []
    for term in _split_terms(raw_value):
        if _is_not_term(term):
            negative_terms.append(term)
    return _dedupe_preserve_order(negative_terms)


def _reduce_terms_with_oncotree(
    terms: Sequence[str],
    hierarchy: OncoTreeHierarchy,
) -> List[str]:
    deduped_terms = _dedupe_preserve_order(terms)
    kept_terms: List[str] = []

    for term in deduped_terms:
        if not hierarchy.has_term(term):
            kept_terms.append(term)
            continue

        should_drop = False
        for other in deduped_terms:
            if other == term:
                continue
            if not hierarchy.has_term(other):
                continue
            if hierarchy.is_more_granular(candidate=other, reference=term):
                should_drop = True
                break

        if not should_drop:
            kept_terms.append(term)

    return _dedupe_preserve_order(kept_terms)


def _resolve_cancer_type_step2(
    cancer_type_step1_raw: object,
    hierarchy: OncoTreeHierarchy,
) -> str:
    all_terms = _dedupe_preserve_order(_split_terms(cancer_type_step1_raw))
    if not all_terms:
        return PAN_CANCER

    positive_terms = _extract_positive_terms(cancer_type_step1_raw)
    negative_terms = _extract_negative_terms(cancer_type_step1_raw)

    if not positive_terms:
        return PAN_CANCER

    reduced_positive_terms = _reduce_terms_with_oncotree(positive_terms, hierarchy)
    combined_terms = _dedupe_preserve_order(reduced_positive_terms + negative_terms)

    if not combined_terms:
        return PAN_CANCER

    return _join_terms(combined_terms)


# ---------------------------------------------------------------------------
# CancerType step 3
# ---------------------------------------------------------------------------

def _should_keep_negative_term_for_step3(
    negative_term: str,
    positive_terms: Sequence[str],
    hierarchy: OncoTreeHierarchy,
) -> bool:
    inner_term = _unwrap_not(negative_term)
    if inner_term is None:
        return False

    # Drop unknown negatives
    if not hierarchy.has_term(inner_term):
        return False

    for positive_term in positive_terms:
        if not hierarchy.has_term(positive_term):
            continue
        if hierarchy.is_more_granular(candidate=inner_term, reference=positive_term):
            return True

    return False


def _filter_negative_terms_for_step3(
    negative_terms: Sequence[str],
    positive_terms: Sequence[str],
    hierarchy: OncoTreeHierarchy,
) -> List[str]:
    kept_negative_terms: List[str] = []

    for negative_term in _dedupe_preserve_order(negative_terms):
        if _should_keep_negative_term_for_step3(
            negative_term,
            positive_terms,
            hierarchy,
        ):
            kept_negative_terms.append(negative_term)

    return kept_negative_terms


def _resolve_cancer_type_step3(
    cancer_type_step2_raw: object,
    hierarchy: OncoTreeHierarchy,
) -> str:
    all_terms = _dedupe_preserve_order(_split_terms(cancer_type_step2_raw))
    if not all_terms:
        return PAN_CANCER

    positive_terms = _extract_positive_terms(cancer_type_step2_raw)
    negative_terms = _extract_negative_terms(cancer_type_step2_raw)

    if not positive_terms:
        return PAN_CANCER

    retained_negative_terms = _filter_negative_terms_for_step3(
        negative_terms,
        positive_terms,
        hierarchy,
    )
    combined_terms = _dedupe_preserve_order(positive_terms + retained_negative_terms)

    if not combined_terms:
        return PAN_CANCER

    return _join_terms(combined_terms)


# ---------------------------------------------------------------------------
# End-to-end workbook transformation
# ---------------------------------------------------------------------------

def append_all_cancer_type_outputs_to_workbook(
    input_excel: Path,
    output_excel: Path,
    curated_dir: Path,
    conditions_mapping_excel: Path,
    oncotree_csv: Path,
    *,
    sheet_name: str = DEFAULT_SHEET_NAME,
    nctid_column_name: str = DEFAULT_NCTID_COLUMN,
    conditions_column_name: str = DEFAULT_CONDITIONS_COLUMN,
    conditions_mapping_sheet_name: Optional[str] = None,
    conditions_lookup_column_name: str = "Conditions_lookup",
    conditions_value_column_name: str = "Oncotree_curation",
    step1_output_column_name: str = DEFAULT_STEP1_COLUMN,
    step2_output_column_name: str = DEFAULT_STEP2_COLUMN,
    step3_output_column_name: str = DEFAULT_STEP3_COLUMN,
) -> Path:
    trial_lookup = build_trial_criteria_lookup(curated_dir)
    conditions_mapping = load_conditions_mapping(
        conditions_mapping_excel,
        sheet_name=conditions_mapping_sheet_name,
        lookup_column_name=conditions_lookup_column_name,
        value_column_name=conditions_value_column_name,
    )
    hierarchy = load_oncotree_hierarchy(oncotree_csv)

    wb = load_workbook(input_excel)
    ws = wb[sheet_name]

    nctid_col_idx = _find_header_column(ws, nctid_column_name)
    conditions_col_idx = _find_header_column(ws, conditions_column_name)

    output_columns = _prepare_output_columns(
        ws,
        [
            "GeneAlteration",
            "GeneAlteration-Inclusive",
            "GeneAlteration-Exclusive",
            "PrimaryTumor",
            "CuratedConditions",
            step1_output_column_name,
            step2_output_column_name,
            step3_output_column_name,
        ],
    )

    nct_rows_seen: Set[str] = set()
    missing_condition_terms: Set[str] = set()

    for row_idx in range(2, ws.max_row + 1):
        raw_nct_id = ws.cell(row=row_idx, column=nctid_col_idx).value
        nct_id = str(raw_nct_id).strip().upper() if raw_nct_id is not None else ""
        if nct_id:
            nct_rows_seen.add(nct_id)

        trial_criteria = trial_lookup.get(nct_id, {})

        gene_alteration_value = _join_cleaned_values(
            "GeneAlteration",
            trial_criteria.get("GeneAlteration", []),
        )

        gene_alteration_inclusive_value, gene_alteration_exclusive_value = (
            _build_gene_alteration_postprocessed(gene_alteration_value)
        )

        primary_tumor_value = _join_cleaned_values(
            "PrimaryTumor",
            trial_criteria.get("PrimaryTumor", []),
        )

        raw_conditions_value = ws.cell(row=row_idx, column=conditions_col_idx).value
        curated_conditions_value = _map_condition_terms(
            raw_conditions_value,
            conditions_mapping,
            missing_condition_terms=missing_condition_terms,
        )

        cancer_type_step1_value = _resolve_cancer_type_step1(
            primary_tumor_value,
            curated_conditions_value,
        )
        cancer_type_step2_value = _resolve_cancer_type_step2(
            cancer_type_step1_value,
            hierarchy,
        )

        cancer_type_step3_value = _resolve_cancer_type_step3(
            cancer_type_step2_value,
            hierarchy,
        )

        ws.cell(row=row_idx, column=output_columns["GeneAlteration"]).value = gene_alteration_value
        ws.cell(
            row=row_idx,
            column=output_columns["GeneAlteration-Inclusive"],
        ).value = gene_alteration_inclusive_value
        ws.cell(
            row=row_idx,
            column=output_columns["GeneAlteration-Exclusive"],
        ).value = gene_alteration_exclusive_value
        ws.cell(row=row_idx, column=output_columns["PrimaryTumor"]).value = primary_tumor_value
        ws.cell(row=row_idx, column=output_columns["CuratedConditions"]).value = curated_conditions_value
        ws.cell(
            row=row_idx,
            column=output_columns[step1_output_column_name],
        ).value = cancer_type_step1_value
        ws.cell(
            row=row_idx,
            column=output_columns[step2_output_column_name],
        ).value = cancer_type_step2_value
        ws.cell(
            row=row_idx,
            column=output_columns[step3_output_column_name],
        ).value = cancer_type_step3_value

    missing_in_excel = sorted(set(trial_lookup) - nct_rows_seen)
    if missing_in_excel:
        logger.warning(
            "%d curated trial(s) were not found in worksheet '%s': %s",
            len(missing_in_excel),
            ws.title,
            ", ".join(missing_in_excel[:20]) + (" ..." if len(missing_in_excel) > 20 else ""),
        )

    if missing_condition_terms:
        logger.warning(
            "%d condition term(s) were not found in mapping '%s': %s",
            len(missing_condition_terms),
            conditions_mapping_excel.name,
            ", ".join(sorted(missing_condition_terms)[:20])
            + (" ..." if len(missing_condition_terms) > 20 else ""),
        )

    output_excel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_excel)
    return output_excel


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Append GeneAlteration, PrimaryTumor, CuratedConditions, "
            "CancerType_step1, CancerType_step2, and CancerType_step3 "
            "to the general worksheet."
        )
    )
    parser.add_argument(
        "--curated_dir",
        required=True,
        type=Path,
        help="Directory containing curated NCT*.py files or a single NCT*.py file",
    )
    parser.add_argument("--input_excel", required=True, type=Path, help="Input Excel workbook")
    parser.add_argument("--output_excel", required=True, type=Path, help="Output Excel workbook")
    parser.add_argument(
        "--conditions_mapping_excel",
        required=True,
        type=Path,
        help="Conditions mapping workbook",
    )
    parser.add_argument(
        "--oncotree_csv",
        required=True,
        type=Path,
        help="OncoTree hierarchy CSV",
    )
    parser.add_argument(
        "--sheet_name",
        default=DEFAULT_SHEET_NAME,
        help="Worksheet to update; defaults to 'general'",
    )
    parser.add_argument(
        "--nctid_column",
        default=DEFAULT_NCTID_COLUMN,
        help="Header name of the NCT ID column",
    )
    parser.add_argument(
        "--conditions_column",
        default=DEFAULT_CONDITIONS_COLUMN,
        help="Header name of the conditions column",
    )
    parser.add_argument(
        "--conditions_mapping_sheet_name",
        default=None,
        help="Worksheet in the conditions mapping workbook; defaults to worksheet 1",
    )
    parser.add_argument(
        "--conditions_lookup_column",
        default="Conditions_lookup",
        help="Lookup column header in the conditions mapping workbook",
    )
    parser.add_argument(
        "--conditions_value_column",
        default="Oncotree_curation",
        help="Return column header in the conditions mapping workbook",
    )
    parser.add_argument(
        "--step1_output_column",
        default=DEFAULT_STEP1_COLUMN,
        help="Header name of the CancerType_step1 column",
    )
    parser.add_argument(
        "--step2_output_column",
        default=DEFAULT_STEP2_COLUMN,
        help="Header name of the CancerType_step2 column",
    )
    parser.add_argument(
        "--step3_output_column",
        default=DEFAULT_STEP3_COLUMN,
        help="Header name of the CancerType_step3 column",
    )
    parser.add_argument("--log_level", default="INFO", help="Logging level (INFO/DEBUG/...)")
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=getattr(logging, str(args.log_level).upper(), logging.INFO),
        format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
    )

    append_all_cancer_type_outputs_to_workbook(
        input_excel=args.input_excel,
        output_excel=args.output_excel,
        curated_dir=args.curated_dir,
        conditions_mapping_excel=args.conditions_mapping_excel,
        oncotree_csv=args.oncotree_csv,
        sheet_name=args.sheet_name,
        nctid_column_name=args.nctid_column,
        conditions_column_name=args.conditions_column,
        conditions_mapping_sheet_name=args.conditions_mapping_sheet_name,
        conditions_lookup_column_name=args.conditions_lookup_column,
        conditions_value_column_name=args.conditions_value_column,
        step1_output_column_name=args.step1_output_column,
        step2_output_column_name=args.step2_output_column,
        step3_output_column_name=args.step3_output_column,
    )
    logger.info("Wrote %s", args.output_excel)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
