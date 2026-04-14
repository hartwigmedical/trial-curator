
from __future__ import annotations

import argparse
import ast
import logging
import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

DELIMITER = " | "
NCT_ID_RE = re.compile(r"(NCT\d{8})", re.IGNORECASE)

CLEANABLE_LABELS = {"PrimaryTumor", "CuratedConditions"}
NULL_SENTINELS = {"[None]", "NOT([None])"}

DEFAULT_SHEET_NAME = "general"
DEFAULT_NCTID_COLUMN = "nctId"
DEFAULT_CONDITIONS_COLUMN = "conditions"

DEFAULT_CANCER_TYPE_LOOKUP_NCTID_COLUMN = "nct_id"
DEFAULT_CANCER_TYPE_INCLUSIVE_COLUMN = "cancer_type_inclusive"
DEFAULT_CANCER_TYPE_EXCLUSIVE_COLUMN = "cancer_type_exclusive"


@dataclass(frozen=True)
class CriterionSpec:
    class_name: str
    label: str
    value_arg_name: str
    split_on_delimiter: bool = False
    emit_inclusive_exclusive: bool = False


CRITERION_SPECS: Dict[str, CriterionSpec] = {
    "GeneAlterationCriterion": CriterionSpec(
        class_name="GeneAlterationCriterion",
        label="GeneAlteration",
        value_arg_name="gene_alteration_curation",
        split_on_delimiter=False,
        emit_inclusive_exclusive=True,
    ),
    "MolecularSignatureCriterion": CriterionSpec(
        class_name="MolecularSignatureCriterion",
        label="MolecularSignature",
        value_arg_name="molecular_signature_curation",
        split_on_delimiter=False,
        emit_inclusive_exclusive=True,
    ),
    "PrimaryTumorCriterion": CriterionSpec(
        class_name="PrimaryTumorCriterion",
        label="PrimaryTumor",
        value_arg_name="Oncotree_curation",
        split_on_delimiter=False,
        emit_inclusive_exclusive=False,
    ),
}

NON_CANCER_OUTPUT_ORDER: List[str] = [
    "GeneAlteration",
    "GeneAlteration-Inclusive",
    "GeneAlteration-Exclusive",
    "MolecularSignature",
    "MolecularSignature-Inclusive",
    "MolecularSignature-Exclusive",
    "PrimaryTumor",
    "CuratedConditions",
]


# ---------------------------------------------------------------------------
# Generic worksheet helpers
# ---------------------------------------------------------------------------

def _find_header_column(ws, header_name: str) -> int:
    for col_idx in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=col_idx).value == header_name:
            return col_idx
    raise ValueError(f"Could not find header '{header_name}' in worksheet '{ws.title}'")


def _copy_cell_style(src, dst) -> None:
    if src.has_style:
        dst._style = copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.font is not None:
        dst.font = copy(src.font)
    if src.fill is not None:
        dst.fill = copy(src.fill)
    if src.border is not None:
        dst.border = copy(src.border)
    if src.alignment is not None:
        dst.alignment = copy(src.alignment)
    if src.protection is not None:
        dst.protection = copy(src.protection)


def _normalize_string(value: object) -> str:
    if value is None:
        return ""
    if pd.isna(value):
        return ""
    return str(value).strip()


def _prepare_output_columns(ws, column_names: Sequence[str]) -> Dict[str, int]:
    col_map: Dict[str, int] = {}

    if ws.max_column < 1:
        raise ValueError(f"Worksheet '{ws.title}' has no columns")

    start_col = ws.max_column + 1
    header_style_template = ws.cell(row=1, column=ws.max_column)
    body_style_template = ws.cell(row=2, column=ws.max_column) if ws.max_row >= 2 else None
    source_width = ws.column_dimensions[get_column_letter(ws.max_column)].width

    for offset, column_name in enumerate(column_names):
        col_idx = start_col + offset
        col_map[column_name] = col_idx

        header_cell = ws.cell(row=1, column=col_idx)
        header_cell.value = column_name
        _copy_cell_style(header_style_template, header_cell)

        if body_style_template is not None:
            for row_idx in range(2, ws.max_row + 1):
                _copy_cell_style(body_style_template, ws.cell(row=row_idx, column=col_idx))

        ws.column_dimensions[get_column_letter(col_idx)].width = max(
            (source_width or 13) * 2.5,
            len(column_name) + 2,
            30,
        )

    return col_map


# ---------------------------------------------------------------------------
# Criteria extraction from curated NCT*.py files
# ---------------------------------------------------------------------------

def _split_top_level_pipe(expr: str) -> List[str]:
    if not expr:
        return []

    parts: List[str] = []
    buf: List[str] = []
    paren = 0
    bracket = 0
    brace = 0
    i = 0
    n = len(expr)

    while i < n:
        ch = expr[i]

        if ch == "(":
            paren += 1
            buf.append(ch)
            i += 1
            continue
        if ch == ")":
            paren = max(paren - 1, 0)
            buf.append(ch)
            i += 1
            continue
        if ch == "[":
            bracket += 1
            buf.append(ch)
            i += 1
            continue
        if ch == "]":
            bracket = max(bracket - 1, 0)
            buf.append(ch)
            i += 1
            continue
        if ch == "{":
            brace += 1
            buf.append(ch)
            i += 1
            continue
        if ch == "}":
            brace = max(brace - 1, 0)
            buf.append(ch)
            i += 1
            continue

        if (
            paren == 0
            and bracket == 0
            and brace == 0
            and expr[i:i + 3] == DELIMITER
        ):
            part = "".join(buf).strip()
            if part:
                parts.append(part)
            buf = []
            i += 3
            continue

        buf.append(ch)
        i += 1

    tail = "".join(buf).strip()
    if tail:
        parts.append(tail)

    return parts


def _is_balanced_wrapped_not(term: str) -> bool:
    if not term:
        return False

    term = term.strip()
    if not term.startswith("NOT(") or not term.endswith(")"):
        return False

    inner = term[4:-1]
    if not inner:
        return False

    paren = 0
    bracket = 0
    brace = 0

    for i, ch in enumerate(term):
        if ch == "(":
            paren += 1
        elif ch == ")":
            paren -= 1
            if paren < 0:
                return False
            if paren == 0 and i != len(term) - 1:
                return False
        elif ch == "[":
            bracket += 1
        elif ch == "]":
            bracket -= 1
            if bracket < 0:
                return False
        elif ch == "{":
            brace += 1
        elif ch == "}":
            brace -= 1
            if brace < 0:
                return False

    return paren == 0 and bracket == 0 and brace == 0


def _unwrap_wrapped_not(term: str) -> str:
    return term.strip()[4:-1].strip()


def _normalize_contradiction_unit(term: str) -> str:
    return re.sub(r"\s+", " ", term.strip())


def _build_inclusive_exclusive_values(raw_value: str) -> Tuple[str, str]:
    if raw_value is None:
        return "", ""

    raw_value = str(raw_value).strip()
    if not raw_value:
        return "", ""

    top_level_terms = _split_top_level_pipe(raw_value)

    positive_terms: List[str] = []
    negative_inner_terms: List[str] = []

    for term in top_level_terms:
        term = term.strip()
        if not term:
            continue

        if _is_balanced_wrapped_not(term):
            negative_inner_terms.append(_unwrap_wrapped_not(term))
        else:
            positive_terms.append(term)

    positive_norms = {_normalize_contradiction_unit(t) for t in positive_terms}
    negative_norms = {_normalize_contradiction_unit(t) for t in negative_inner_terms}
    contradictions = positive_norms & negative_norms

    if contradictions:
        positive_terms = [
            t for t in positive_terms
            if _normalize_contradiction_unit(t) not in contradictions
        ]
        negative_inner_terms = [
            t for t in negative_inner_terms
            if _normalize_contradiction_unit(t) not in contradictions
        ]

    inclusive_value = DELIMITER.join(positive_terms) if positive_terms else ""
    exclusive_value = f"NOT({DELIMITER.join(negative_inner_terms)})" if negative_inner_terms else ""

    return inclusive_value, exclusive_value


def _iter_input_py_files(curated_dir: Path) -> Iterable[Path]:
    if curated_dir.is_file():
        yield curated_dir
        return
    yield from sorted(p for p in curated_dir.glob("NCT*.py") if p.is_file())


def _extract_nct_id_from_path(py_path: Path) -> str:
    match = NCT_ID_RE.search(py_path.stem)
    if not match:
        raise ValueError(f"Could not extract NCT ID from filename: {py_path.name}")
    return match.group(1).upper()


def _call_name(node: ast.AST) -> Optional[str]:
    if isinstance(node, ast.Name):
        return node.id
    if isinstance(node, ast.Attribute):
        return node.attr
    return None


def _get_call_argument(call: ast.Call, key: str) -> Optional[ast.AST]:
    for kw in call.keywords:
        if kw.arg == key:
            return kw.value
    return None


def _get_string_constant(node: Optional[ast.AST]) -> Optional[str]:
    if isinstance(node, ast.Constant) and isinstance(node.value, str):
        return node.value
    return None


def _split_emitted_values(raw_value: str, split_on_delimiter: bool) -> List[str]:
    if split_on_delimiter:
        return [part.strip() for part in raw_value.split(DELIMITER) if part.strip()]
    stripped = raw_value.strip()
    return [stripped] if stripped else []


def _apply_not_wrapper(value: str, is_negated: bool) -> str:
    if not is_negated:
        return value
    if value.startswith("NOT("):
        return value
    return f"NOT({value})"


def _clean_value(criterion_label: str, value: str) -> Optional[str]:
    stripped = value.strip()
    if criterion_label in CLEANABLE_LABELS and stripped in NULL_SENTINELS:
        return None
    return stripped


def _finalize_value(criterion_label: str, value: str, *, is_negated: bool) -> Optional[str]:
    wrapped = _apply_not_wrapper(value.strip(), is_negated=is_negated)
    return _clean_value(criterion_label, wrapped)


def _add_extracted_value(
    bucket: Dict[str, List[str]],
    criterion_label: str,
    raw_value: str,
    *,
    is_negated: bool,
    split_on_delimiter: bool,
) -> None:
    parts = _split_emitted_values(raw_value, split_on_delimiter=split_on_delimiter)
    if not parts:
        return

    for part in parts:
        final_value = _finalize_value(
            criterion_label,
            part,
            is_negated=is_negated,
        )
        if final_value is not None:
            bucket.setdefault(criterion_label, []).append(final_value)


def _walk_ast(node: ast.AST, bucket: Dict[str, List[str]], *, negation_depth: int = 0) -> None:
    if isinstance(node, ast.Call):
        func_name = _call_name(node.func)

        if func_name == "NotCriterion":
            child = _get_call_argument(node, "criterion")
            if child is None and node.args:
                child = node.args[0]
            if child is not None:
                _walk_ast(child, bucket, negation_depth=negation_depth + 1)
            return

        target_spec = CRITERION_SPECS.get(func_name)
        if target_spec is not None:
            raw_value = _get_string_constant(_get_call_argument(node, target_spec.value_arg_name))
            if raw_value:
                _add_extracted_value(
                    bucket,
                    target_spec.label,
                    raw_value,
                    is_negated=negation_depth > 0,
                    split_on_delimiter=target_spec.split_on_delimiter,
                )

    for child in ast.iter_child_nodes(node):
        _walk_ast(child, bucket, negation_depth=negation_depth)


def _collapse_primary_tumor_contradictions(values: Sequence[str]) -> List[str]:
    value_set = set(values)
    collapsed: Set[str] = set()

    for value in value_set:
        if value.startswith("NOT(") and value.endswith(")"):
            positive_value = value[4:-1]
            if positive_value in value_set:
                continue
        collapsed.add(value)

    return sorted(collapsed)


def _dedupe_and_sort(criterion_label: str, values: Sequence[str]) -> List[str]:
    deduped = {value for value in values if value}
    if criterion_label == "PrimaryTumor":
        return _collapse_primary_tumor_contradictions(deduped)
    return sorted(deduped)


def _join_cleaned_values(criterion_label: str, values: Sequence[str]) -> str:
    return DELIMITER.join(_dedupe_and_sort(criterion_label, values))


def extract_criteria_from_py(py_path: Path) -> Dict[str, List[str]]:
    tree = ast.parse(py_path.read_text(encoding="utf-8"), filename=str(py_path))
    extracted: Dict[str, List[str]] = {}
    _walk_ast(tree, extracted)
    return {
        criterion: _dedupe_and_sort(criterion, values)
        for criterion, values in extracted.items()
        if values
    }


def build_trial_criteria_lookup(curated_dir: Path) -> Dict[str, Dict[str, List[str]]]:
    lookup: Dict[str, Dict[str, List[str]]] = {}
    seen_paths: Dict[str, Path] = {}

    for py_path in _iter_input_py_files(curated_dir):
        nct_id = _extract_nct_id_from_path(py_path)
        if nct_id in lookup:
            raise ValueError(
                f"Duplicate curated .py files for {nct_id}: {seen_paths[nct_id]} and {py_path}"
            )
        seen_paths[nct_id] = py_path
        lookup[nct_id] = extract_criteria_from_py(py_path)

    return lookup


# ---------------------------------------------------------------------------
# Conditions mapping
# ---------------------------------------------------------------------------

def _parse_condition_terms(raw_value: object) -> List[str]:
    if raw_value is None:
        return []

    if isinstance(raw_value, list):
        return [str(item).strip() for item in raw_value if str(item).strip()]

    text = str(raw_value).strip()
    if not text:
        return []

    try:
        parsed = ast.literal_eval(text)
    except (SyntaxError, ValueError):
        return [text]

    if isinstance(parsed, list):
        return [str(item).strip() for item in parsed if str(item).strip()]

    return [text]


def _map_condition_terms(
    raw_value: object,
    conditions_mapping: Dict[str, str],
    *,
    missing_condition_terms: Set[str],
) -> str:
    mapped_terms: List[str] = []

    for term in _parse_condition_terms(raw_value):
        mapped_term = _normalize_string(conditions_mapping.get(term, ""))
        if mapped_term:
            cleaned = _clean_value("CuratedConditions", mapped_term)
            if cleaned is not None:
                mapped_terms.append(cleaned)
        else:
            missing_condition_terms.add(term)

    return _join_cleaned_values("CuratedConditions", mapped_terms)


def load_conditions_mapping(
    mapping_excel: Path,
    *,
    sheet_name: Optional[str] = None,
    lookup_column_name: str = "Conditions_lookup",
    value_column_name: str = "Oncotree_curation",
) -> Dict[str, str]:
    wb = load_workbook(mapping_excel, read_only=True, data_only=True)
    target_sheet_name = sheet_name or wb.sheetnames[0]
    ws = wb[target_sheet_name]

    lookup_col_idx = _find_header_column(ws, lookup_column_name)
    value_col_idx = _find_header_column(ws, value_column_name)

    mapping: Dict[str, str] = {}
    for row_idx in range(2, ws.max_row + 1):
        lookup_value = _normalize_string(ws.cell(row=row_idx, column=lookup_col_idx).value)
        curated_value = _normalize_string(ws.cell(row=row_idx, column=value_col_idx).value)

        if not lookup_value:
            continue
        mapping[lookup_value] = curated_value

    return mapping


# ---------------------------------------------------------------------------
# Cancer type lookup
# ---------------------------------------------------------------------------

def load_cancer_type_lookup(
    cancer_types_file: Path,
    *,
    nctid_column_name: str = DEFAULT_CANCER_TYPE_LOOKUP_NCTID_COLUMN,
    inclusive_column_name: str = DEFAULT_CANCER_TYPE_INCLUSIVE_COLUMN,
    exclusive_column_name: str = DEFAULT_CANCER_TYPE_EXCLUSIVE_COLUMN,
) -> Dict[str, Dict[str, str]]:
    suffix = cancer_types_file.suffix.lower()
    if suffix == ".csv":
        df = pd.read_csv(cancer_types_file, dtype=str, keep_default_na=False, na_values=[])
    elif suffix in {".xlsx", ".xls"}:
        df = pd.read_excel(cancer_types_file, dtype=str, keep_default_na=False, na_values=[])
    else:
        raise ValueError(f"Unsupported cancer types file type: {cancer_types_file.suffix}")

    df.columns = [str(c).strip() for c in df.columns]

    required = {nctid_column_name, inclusive_column_name, exclusive_column_name}
    missing = [col for col in required if col not in df.columns]
    if missing:
        raise ValueError(
            f"Cancer types file missing required column(s): {missing}. "
            f"Found columns: {list(df.columns)}"
        )

    lookup: Dict[str, Dict[str, str]] = {}
    for _, row in df.iterrows():
        nct_id = _normalize_string(row.get(nctid_column_name, "")).upper()
        if not nct_id:
            continue
        lookup[nct_id] = {
            inclusive_column_name: _normalize_string(row.get(inclusive_column_name, "")),
            exclusive_column_name: _normalize_string(row.get(exclusive_column_name, "")),
        }

    return lookup


# ---------------------------------------------------------------------------
# Criterion output shaping
# ---------------------------------------------------------------------------

def _build_trial_output_values(
    trial_criteria: Dict[str, List[str]],
) -> Dict[str, str]:
    out: Dict[str, str] = {}

    for spec in CRITERION_SPECS.values():
        raw_value = _join_cleaned_values(
            spec.label,
            trial_criteria.get(spec.label, []),
        )
        out[spec.label] = raw_value

        if spec.emit_inclusive_exclusive:
            inclusive_value, exclusive_value = _build_inclusive_exclusive_values(raw_value)
            out[f"{spec.label}-Inclusive"] = inclusive_value
            out[f"{spec.label}-Exclusive"] = exclusive_value

    return out


# ---------------------------------------------------------------------------
# End-to-end workbook transformation
# ---------------------------------------------------------------------------

def append_trial_outputs_to_workbook(
    extraction_excel: Path,
    output_excel: Path,
    curated_dir: Path,
    conditions_mapping_excel: Path,
    cancer_types_file: Path,
    *,
    sheet_name: str = DEFAULT_SHEET_NAME,
    nctid_column_name: str = DEFAULT_NCTID_COLUMN,
    conditions_column_name: str = DEFAULT_CONDITIONS_COLUMN,
    conditions_mapping_sheet_name: Optional[str] = None,
    conditions_lookup_column_name: str = "Conditions_lookup",
    conditions_value_column_name: str = "Oncotree_curation",
    cancer_type_lookup_nctid_column: str = DEFAULT_CANCER_TYPE_LOOKUP_NCTID_COLUMN,
    cancer_type_inclusive_column: str = DEFAULT_CANCER_TYPE_INCLUSIVE_COLUMN,
    cancer_type_exclusive_column: str = DEFAULT_CANCER_TYPE_EXCLUSIVE_COLUMN,
) -> Path:
    trial_lookup = build_trial_criteria_lookup(curated_dir)
    conditions_mapping = load_conditions_mapping(
        conditions_mapping_excel,
        sheet_name=conditions_mapping_sheet_name,
        lookup_column_name=conditions_lookup_column_name,
        value_column_name=conditions_value_column_name,
    )
    cancer_type_lookup = load_cancer_type_lookup(
        cancer_types_file,
        nctid_column_name=cancer_type_lookup_nctid_column,
        inclusive_column_name=cancer_type_inclusive_column,
        exclusive_column_name=cancer_type_exclusive_column,
    )

    wb = load_workbook(extraction_excel)
    ws = wb[sheet_name]

    nctid_col_idx = _find_header_column(ws, nctid_column_name)
    conditions_col_idx = _find_header_column(ws, conditions_column_name)

    output_columns = _prepare_output_columns(
        ws,
        NON_CANCER_OUTPUT_ORDER
        + [
            cancer_type_inclusive_column,
            cancer_type_exclusive_column,
        ],
    )

    nct_rows_seen: Set[str] = set()
    missing_condition_terms: Set[str] = set()
    missing_cancer_type_trials: Set[str] = set()

    for row_idx in range(2, ws.max_row + 1):
        raw_nct_id = ws.cell(row=row_idx, column=nctid_col_idx).value
        nct_id = _normalize_string(raw_nct_id).upper()
        if nct_id:
            nct_rows_seen.add(nct_id)

        trial_criteria = trial_lookup.get(nct_id, {})
        criterion_outputs = _build_trial_output_values(trial_criteria)

        raw_conditions_value = ws.cell(row=row_idx, column=conditions_col_idx).value
        curated_conditions_value = _map_condition_terms(
            raw_conditions_value,
            conditions_mapping,
            missing_condition_terms=missing_condition_terms,
        )
        criterion_outputs["CuratedConditions"] = curated_conditions_value

        cancer_type_values = cancer_type_lookup.get(nct_id, {})
        if nct_id and not cancer_type_values:
            missing_cancer_type_trials.add(nct_id)

        ws.cell(row=row_idx, column=output_columns["GeneAlteration"]).value = criterion_outputs.get("GeneAlteration", "")
        ws.cell(row=row_idx, column=output_columns["GeneAlteration-Inclusive"]).value = criterion_outputs.get("GeneAlteration-Inclusive", "")
        ws.cell(row=row_idx, column=output_columns["GeneAlteration-Exclusive"]).value = criterion_outputs.get("GeneAlteration-Exclusive", "")
        ws.cell(row=row_idx, column=output_columns["MolecularSignature"]).value = criterion_outputs.get("MolecularSignature", "")
        ws.cell(row=row_idx, column=output_columns["MolecularSignature-Inclusive"]).value = criterion_outputs.get("MolecularSignature-Inclusive", "")
        ws.cell(row=row_idx, column=output_columns["MolecularSignature-Exclusive"]).value = criterion_outputs.get("MolecularSignature-Exclusive", "")
        ws.cell(row=row_idx, column=output_columns["PrimaryTumor"]).value = criterion_outputs.get("PrimaryTumor", "")
        ws.cell(row=row_idx, column=output_columns["CuratedConditions"]).value = curated_conditions_value
        ws.cell(row=row_idx, column=output_columns[cancer_type_inclusive_column]).value = cancer_type_values.get(cancer_type_inclusive_column, "")
        ws.cell(row=row_idx, column=output_columns[cancer_type_exclusive_column]).value = cancer_type_values.get(cancer_type_exclusive_column, "")

    missing_in_excel = sorted(set(trial_lookup) - nct_rows_seen)
    if missing_in_excel:
        logger.warning(
            "%d curated trial(s) were not found in worksheet '%s': %s",
            len(missing_in_excel),
            ws.title,
            ", ".join(missing_in_excel[:20]) + (" ..." if len(missing_in_excel) > 20 else ""),
        )

    if missing_condition_terms:
        logger.warning(
            "%d condition term(s) were not found in mapping '%s': %s",
            len(missing_condition_terms),
            conditions_mapping_excel.name,
            ", ".join(sorted(missing_condition_terms)[:20])
            + (" ..." if len(missing_condition_terms) > 20 else ""),
        )

    if missing_cancer_type_trials:
        logger.warning(
            "%d trial(s) in the workbook were not found in cancer types file '%s': %s",
            len(missing_cancer_type_trials),
            cancer_types_file.name,
            ", ".join(sorted(missing_cancer_type_trials)[:20])
            + (" ..." if len(missing_cancer_type_trials) > 20 else ""),
        )

    output_excel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_excel)
    return output_excel


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Append GeneAlteration, MolecularSignature, PrimaryTumor, CuratedConditions, "
            "and cancer_type_inclusive/cancer_type_exclusive to the general worksheet."
        )
    )
    parser.add_argument(
        "--curated_dir",
        required=True,
        type=Path,
        help="Directory containing curated NCT*.py files or a single NCT*.py file",
    )
    parser.add_argument("--extraction_excel", required=True, type=Path, help="Input Excel workbook")
    parser.add_argument("--output_excel", required=True, type=Path, help="Output Excel workbook")
    parser.add_argument(
        "--conditions_mapping_excel",
        required=True,
        type=Path,
        help="Conditions mapping workbook",
    )
    parser.add_argument(
        "--cancer_types_file",
        required=True,
        type=Path,
        help="CSV/XLSX containing nct_id, cancer_type_inclusive, and cancer_type_exclusive",
    )
    parser.add_argument(
        "--sheet_name",
        default=DEFAULT_SHEET_NAME,
        help="Worksheet to update; defaults to 'general'",
    )
    parser.add_argument(
        "--nctid_column",
        default=DEFAULT_NCTID_COLUMN,
        help="Header name of the NCT ID column in the workbook",
    )
    parser.add_argument(
        "--conditions_column",
        default=DEFAULT_CONDITIONS_COLUMN,
        help="Header name of the conditions column in the workbook",
    )
    parser.add_argument(
        "--conditions_mapping_sheet_name",
        default=None,
        help="Worksheet in the conditions mapping workbook; defaults to worksheet 1",
    )
    parser.add_argument(
        "--conditions_lookup_column",
        default="Conditions_lookup",
        help="Lookup column header in the conditions mapping workbook",
    )
    parser.add_argument(
        "--conditions_value_column",
        default="Oncotree_curation",
        help="Return column header in the conditions mapping workbook",
    )
    parser.add_argument(
        "--cancer_type_lookup_nctid_column",
        default=DEFAULT_CANCER_TYPE_LOOKUP_NCTID_COLUMN,
        help="NCT ID column header in the cancer types file",
    )
    parser.add_argument(
        "--cancer_type_inclusive_column",
        default=DEFAULT_CANCER_TYPE_INCLUSIVE_COLUMN,
        help="Inclusive cancer type column header in the cancer types file and output workbook",
    )
    parser.add_argument(
        "--cancer_type_exclusive_column",
        default=DEFAULT_CANCER_TYPE_EXCLUSIVE_COLUMN,
        help="Exclusive cancer type column header in the cancer types file and output workbook",
    )
    parser.add_argument("--log_level", default="INFO", help="Logging level (INFO/DEBUG/...)")
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=getattr(logging, str(args.log_level).upper(), logging.INFO),
        format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
    )

    append_trial_outputs_to_workbook(
        extraction_excel=args.extraction_excel,
        output_excel=args.output_excel,
        curated_dir=args.curated_dir,
        conditions_mapping_excel=args.conditions_mapping_excel,
        cancer_types_file=args.cancer_types_file,
        sheet_name=args.sheet_name,
        nctid_column_name=args.nctid_column,
        conditions_column_name=args.conditions_column,
        conditions_mapping_sheet_name=args.conditions_mapping_sheet_name,
        conditions_lookup_column_name=args.conditions_lookup_column,
        conditions_value_column_name=args.conditions_value_column,
        cancer_type_lookup_nctid_column=args.cancer_type_lookup_nctid_column,
        cancer_type_inclusive_column=args.cancer_type_inclusive_column,
        cancer_type_exclusive_column=args.cancer_type_exclusive_column,
    )
    logger.info("Wrote %s", args.output_excel)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
